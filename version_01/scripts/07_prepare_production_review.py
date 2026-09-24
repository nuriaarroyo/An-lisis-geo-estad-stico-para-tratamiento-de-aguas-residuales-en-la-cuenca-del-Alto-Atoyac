from __future__ import annotations

from pathlib import Path
from urllib.parse import quote_plus

import geopandas as gpd
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from version_01.scripts.common import PROJECT_ROOT, normalize_text


INPUT_PATH = PROJECT_ROOT / "outputs/auditoria_6_especiales/capas_qgis/universo_productivo_auditado.gpkg"
OUT_DIR = PROJECT_ROOT / "outputs/revision_produccion"
MH_PATH = PROJECT_ROOT / "data/processed/capa_maryhelen/capa MH.shp"


def clean(value: object) -> str:
    return "" if pd.isna(value) else str(value).strip()


def full_address(row: pd.Series) -> str:
    street = " ".join(value for value in [clean(row.get("tipo_vial")), clean(row.get("nom_vial")), clean(row.get("numero_ext")), clean(row.get("letra_ext"))] if value)
    settlement = " ".join(value for value in [clean(row.get("tipo_asent")), clean(row.get("nomb_asent"))] if value)
    return ", ".join(value for value in [street, settlement, clean(row.get("cod_postal")), clean(row.get("localidad")), clean(row.get("municipio")), clean(row.get("entidad"))] if value)


def maps_search(row: pd.Series) -> str:
    query = " ".join(value for value in [clean(row.get("tipo_vial")), clean(row.get("nom_vial")), clean(row.get("numero_ext")), clean(row.get("letra_ext")), clean(row.get("tipo_asent")), clean(row.get("nomb_asent")), clean(row.get("cod_postal")), clean(row.get("localidad")), clean(row.get("municipio")), clean(row.get("entidad"))] if value)
    return "https://www.google.com/maps/search/?api=1&query=" + quote_plus(query)


def format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    links = {
        "url_busqueda_google_maps": "Abrir búsqueda en Google Maps",
        "url_punto_google_maps": "Abrir punto en Google Maps",
        "url_openstreetmap": "Abrir punto en OpenStreetMap",
    }
    for name, label in links.items():
        for row_number in range(2, sheet.max_row + 1):
            cell = sheet.cell(row_number, headers[name])
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.value = label
                cell.font = Font(color="0563C1", underline="single")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def main() -> None:
    if not INPUT_PATH.exists():
        raise FileNotFoundError("Primero ejecute 05_apply_special_audit.py")
    universe = gpd.read_file(INPUT_PATH, engine="pyogrio").to_crs("EPSG:4326")
    production = universe.loc[universe["categoria_clasificacion"].eq("produccion_textil")].copy()
    mh = gpd.read_file(MH_PATH, engine="pyogrio")
    mh_ids = set(mh["ID DENUE"].fillna("").astype(str).str.replace(r"\.0$", "", regex=True).str.strip()) - {"", "nan", "None"}
    production["en_mh"] = production["id"].astype(str).str.replace(r"\.0$", "", regex=True).isin(mh_ids)
    production["antecedente_mh"] = "NO_APARECE_EN_MH"
    production.loc[production["en_mh"], "antecedente_mh"] = "VISTO_BUENO_PREVIO_MH"
    production["nota_antecedente_mh"] = ""
    production.loc[production["en_mh"], "nota_antecedente_mh"] = "MH incluyó previamente este ID DENUE; usar como antecedente favorable, sin asumir que confirma proceso húmedo actual"
    names = production["nombre_de_la_unidad_economica"].fillna("").map(normalize_text)

    wet_pattern = r"lavado|lavander|deslav|deslave|stone wash|stone lav|acid wash|tenid|tintura|blanque|mercer|enzim|sosa caustica|colorante|bano quimico|quimic|suaviz|pigment|decolor|neutraliz|centrifug|enjuag|oxidante"
    denim_pattern = r"mezclilla|mesclilla|denim|jeans?"
    dry_pattern = r"\bcorte\b|costur|confeccion|deshebr|desebrad|bordad|\bover\b|\brecta\b|hojale|ensambl|presilla|etiquet|dobladill|armado"
    indirect_dry_pattern = r"planch"
    ambiguous_pattern = r"acabad|terminado|preparado|elaboracion|transformacion"
    wet = names.str.contains(wet_pattern, regex=True)
    denim = names.str.contains(denim_pattern, regex=True)
    dry = names.str.contains(dry_pattern, regex=True)
    indirect_dry = names.str.contains(indirect_dry_pattern, regex=True)
    ambiguous = names.str.contains(ambiguous_pattern, regex=True)

    production["clasificacion_proceso_hidrico"] = "PRODUCCION_SIN_PROCESO_DECLARADO"
    production["decision_hidrica_automatica"] = "REVISAR"
    production["prioridad_hidrica_produccion"] = "revision"
    production["causa_decision_hidrica"] = "DENUE confirma producción textil, pero el nombre no permite determinar si el proceso usa agua"
    production.loc[ambiguous, ["clasificacion_proceso_hidrico", "causa_decision_hidrica"]] = ["ACABADO_PREPARADO_AMBIGUO", "El nombre indica acabado, terminado, preparado o elaboración sin describir el proceso"]
    production.loc[dry | indirect_dry, ["clasificacion_proceso_hidrico", "decision_hidrica_automatica", "prioridad_hidrica_produccion"]] = ["SECO_EXPLICITO", "EXCLUIR_SECO", "fuera_hidrico"]
    production.loc[dry, "causa_decision_hidrica"] = "El nombre declara corte, costura, confección, deshebrado, bordado o ensamble; proceso seco fuera del universo hídrico"
    production.loc[indirect_dry, "causa_decision_hidrica"] = "El nombre declara planchado; consumo indirecto insuficiente para prioridad hídrica"
    production.loc[denim, ["clasificacion_proceso_hidrico", "decision_hidrica_automatica", "prioridad_hidrica_produccion", "causa_decision_hidrica"]] = ["MEZCLILLA_DENIM_EXPLICITA", "INCLUIR_POR_MEZCLILLA", "inclusion_precautoria", "El nombre menciona mezclilla, mesclilla, denim o jeans; inclusión precautoria sin asumir que exista lavado"]
    production.loc[wet, ["clasificacion_proceso_hidrico", "decision_hidrica_automatica", "prioridad_hidrica_produccion", "causa_decision_hidrica"]] = ["HUMEDO_EXPLICITO", "INCLUIR_PROCESO_HUMEDO", "alta", "El nombre declara una señal fuerte de proceso húmedo o químico textil"]

    production["latitud_revision"] = production.geometry.y
    production["longitud_revision"] = production.geometry.x
    production["direccion_completa"] = production.apply(full_address, axis=1)
    production["url_busqueda_google_maps"] = production.apply(maps_search, axis=1)
    production["url_punto_google_maps"] = production.apply(lambda row: "https://www.google.com/maps/search/?api=1&query=" + quote_plus(f"{row['latitud_revision']},{row['longitud_revision']}"), axis=1)
    production["url_openstreetmap"] = production.apply(lambda row: f"https://www.openstreetmap.org/?mlat={row['latitud_revision']:.7f}&mlon={row['longitud_revision']:.7f}#map=19/{row['latitud_revision']:.7f}/{row['longitud_revision']:.7f}", axis=1)
    production["estado_revision_manual"] = "PENDIENTE"
    production.loc[production["decision_hidrica_automatica"].ne("REVISAR"), "estado_revision_manual"] = "RESUELTO_AUTOMATICAMENTE"
    for column in ["establecimiento_encontrado", "actividad_observada", "evidencia_textil", "evidencia_mezclilla", "evidencia_proceso_humedo", "decision_manual", "confianza_revision", "fuentes_consultadas", "fecha_revision", "revisor", "notas_revision"]:
        production[column] = ""

    columns = [
        "id", "clee", "nombre_de_la_unidad_economica", "raz_social", "codigo_de_la_clase_actividad_scian", "nombre_clase_actividad_scian", "categoria_clasificacion", "regla_decisiva", "terminos_detectados", "per_ocu",
        "clasificacion_proceso_hidrico", "decision_hidrica_automatica", "prioridad_hidrica_produccion", "causa_decision_hidrica", "en_mh", "antecedente_mh", "nota_antecedente_mh", "direccion_completa", "tipo_vial", "nom_vial", "numero_ext", "letra_ext", "numero_int", "letra_int", "tipo_asent", "nomb_asent", "cod_postal", "localidad", "municipio", "entidad", "telefono", "correoelec", "www", "latitud_revision", "longitud_revision", "url_busqueda_google_maps", "url_punto_google_maps", "url_openstreetmap",
        "estado_revision_manual", "establecimiento_encontrado", "actividad_observada", "evidencia_textil", "evidencia_mezclilla", "evidencia_proceso_humedo", "decision_manual", "confianza_revision", "fuentes_consultadas", "fecha_revision", "revisor", "notas_revision",
    ]
    order = {"INCLUIR_PROCESO_HUMEDO": 0, "INCLUIR_POR_MEZCLILLA": 1, "REVISAR": 2, "EXCLUIR_SECO": 9}
    production["_orden"] = production["decision_hidrica_automatica"].map(order)
    queue = production.sort_values(["_orden", "nombre_de_la_unidad_economica", "id"])[columns]
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    queue.to_csv(OUT_DIR / "cola_revision_produccion.csv", index=False, encoding="utf-8-sig")
    pending_review = queue.loc[queue["decision_hidrica_automatica"].eq("REVISAR")]
    pending_xlsx = OUT_DIR / "cola_revision_produccion_pendiente.xlsx"
    try:
        pending_review.to_excel(pending_xlsx, index=False, engine="openpyxl")
    except PermissionError:
        pending_xlsx = OUT_DIR / "cola_revision_produccion_pendiente_actualizada.xlsx"
        pending_review.to_excel(pending_xlsx, index=False, engine="openpyxl")
    format_workbook(pending_xlsx)
    try:
        pending_review.to_csv(OUT_DIR / "cola_revision_produccion_pendiente.csv", index=False, encoding="utf-8-sig")
    except PermissionError:
        pending_review.to_csv(OUT_DIR / "cola_revision_produccion_pendiente_actualizada.csv", index=False, encoding="utf-8-sig")
    queue.loc[queue["decision_hidrica_automatica"].str.startswith("INCLUIR")].to_csv(OUT_DIR / "produccion_inclusion_automatica_hidrica.csv", index=False, encoding="utf-8-sig")
    queue.loc[queue["decision_hidrica_automatica"].eq("EXCLUIR_SECO")].to_csv(OUT_DIR / "produccion_excluida_proceso_seco.csv", index=False, encoding="utf-8-sig")
    queue.groupby(["clasificacion_proceso_hidrico", "decision_hidrica_automatica"], dropna=False).size().reset_index(name="registros").to_csv(OUT_DIR / "resumen_revision_produccion.csv", index=False, encoding="utf-8-sig")
    xlsx_path = OUT_DIR / "cola_revision_produccion.xlsx"
    try:
        queue.to_excel(xlsx_path, index=False, engine="openpyxl")
    except PermissionError:
        xlsx_path = OUT_DIR / "cola_revision_produccion_actualizada.xlsx"
        queue.to_excel(xlsx_path, index=False, engine="openpyxl")
    format_workbook(xlsx_path)
    print(f"Cola de producción preparada: {len(queue)} registros")


if __name__ == "__main__":
    main()
