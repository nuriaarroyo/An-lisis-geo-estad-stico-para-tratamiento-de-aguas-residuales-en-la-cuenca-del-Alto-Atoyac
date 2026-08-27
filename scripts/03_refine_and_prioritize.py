from __future__ import annotations

import hashlib
import json
from datetime import datetime
from pathlib import Path
from urllib.parse import quote_plus

import geopandas as gpd
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from common import PROJECT_ROOT, normalize_text, safe_to_file


CONFIG_DIR = PROJECT_ROOT / "config" / "filtros_textiles_santa_ana" / "versiones" / "v1"
STAGE1_DIR = PROJECT_ROOT / "outputs" / "universo_independiente_3_localidades"
STAGE1_INPUT = STAGE1_DIR / "capas_qgis" / "universo_independiente_incluido.gpkg"
STAGE1_REVIEW = STAGE1_DIR / "capas_qgis" / "universo_independiente_revision.gpkg"
STAGE1_MANIFEST = STAGE1_DIR / "manifest_clasificacion.json"
OUT_DIR = PROJECT_ROOT / "outputs" / "universo_refinado_santa_ana"
REFERENCES_CONFIG = CONFIG_DIR / "referencias_comparacion.csv"


def flag_set(value: object) -> set[str]:
    if pd.isna(value):
        return set()
    return {part.strip() for part in str(value).split("|") if part.strip()}


def choose_rule(flags: dict[str, bool], policy: pd.DataFrame) -> pd.Series:
    for _, rule in policy.sort_values("priority", ascending=False).iterrows():
        required = flag_set(rule.get("all_flags", ""))
        alternatives = flag_set(rule.get("any_flags", ""))
        forbidden = flag_set(rule.get("none_flags", ""))
        if required and not all(flags.get(flag, False) for flag in required):
            continue
        if alternatives and not any(flags.get(flag, False) for flag in alternatives):
            continue
        if forbidden and any(flags.get(flag, False) for flag in forbidden):
            continue
        return rule
    raise ValueError("La política de refinamiento no contiene una regla aplicable")


def export(frame: gpd.GeoDataFrame, name: str) -> None:
    safe_to_file(frame, OUT_DIR / "capas_qgis" / f"{name}.gpkg", layer=name)
    pd.DataFrame(frame.drop(columns="geometry")).to_csv(
        OUT_DIR / "tablas" / f"{name}.csv", index=False, encoding="utf-8-sig"
    )


def normalize_id(series: pd.Series) -> pd.Series:
    return series.fillna("").astype(str).str.replace(r"\.0$", "", regex=True).str.strip().str.upper()


def load_mh_ids() -> set[str]:
    references = pd.read_csv(REFERENCES_CONFIG, dtype=str).fillna("")
    row = references.loc[references["referencia_id"].eq("mh_filtrado")].iloc[0]
    mh = gpd.read_file(PROJECT_ROOT / row["ruta"], engine="pyogrio")
    return set(normalize_id(mh[row["campo_id"]])) - {""}


def explicit_review_cause(category: str) -> str:
    return {
        "revisar_lavado_generico": (
            "SCIAN 812210 y nombre de lavandería sin evidencia explícita de lavado industrial, "
            "mezclilla, deslavado o acabado productivo"
        ),
        "revisar_contradiccion_scian_texto": (
            "SCIAN 812210 contradicho por nombre de bodega o desperdicio"
        ),
        "revisar_maquila_sola": (
            "Nombre de maquila sin evidencia de objeto textil y con SCIAN no manufacturero"
        ),
    }.get(category, "Señal insuficiente; requiere revisión humana")


def clean_value(value: object) -> str:
    return "" if pd.isna(value) else str(value).strip()


def format_review_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    links = {
        "url_busqueda_google_maps": "Buscar dirección DENUE en Google Maps",
        "url_punto_google_maps": "Abrir punto DENUE en Google Maps",
        "url_openstreetmap": "Abrir punto en OpenStreetMap",
    }
    for name, label in links.items():
        if name not in headers:
            continue
        for row_number in range(2, sheet.max_row + 1):
            cell = sheet.cell(row_number, headers[name])
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.value = label
                cell.font = Font(color="0563C1", underline="single")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def manual_review_queue(frame: gpd.GeoDataFrame, universe: str) -> pd.DataFrame:
    geographic = frame.to_crs("EPSG:4326").copy()
    geographic["latitud_revision"] = geographic.geometry.y
    geographic["longitud_revision"] = geographic.geometry.x

    def search_url(row: pd.Series) -> str:
        parts = [
            clean_value(row.get(column)) for column in [
                "tipo_vial", "nom_vial", "numero_ext", "letra_ext",
                "tipo_asent", "nomb_asent", "cod_postal", "localidad",
                "municipio", "entidad",
            ]
        ]
        parts = [value for value in parts if value]
        return "https://www.google.com/maps/search/?api=1&query=" + quote_plus(" ".join(parts))

    geographic["direccion_completa_denue"] = geographic.apply(
        lambda row: ", ".join(
            value for value in [
                " ".join(value for value in [clean_value(row.get("tipo_vial")), clean_value(row.get("nom_vial")), clean_value(row.get("numero_ext")), clean_value(row.get("letra_ext"))] if value),
                " ".join(value for value in [clean_value(row.get("tipo_asent")), clean_value(row.get("nomb_asent"))] if value),
                clean_value(row.get("cod_postal")), clean_value(row.get("localidad")),
                clean_value(row.get("municipio")), clean_value(row.get("entidad")),
            ] if value
        ), axis=1,
    )
    geographic["url_busqueda_google_maps"] = geographic.apply(search_url, axis=1)
    geographic["url_punto_google_maps"] = geographic.apply(
        lambda row: "https://www.google.com/maps/search/?api=1&query="
        + quote_plus(f"{row['latitud_revision']},{row['longitud_revision']}"),
        axis=1,
    )
    geographic["url_openstreetmap"] = geographic.apply(
        lambda row: (
            f"https://www.openstreetmap.org/?mlat={row['latitud_revision']:.7f}"
            f"&mlon={row['longitud_revision']:.7f}#map=19/"
            f"{row['latitud_revision']:.7f}/{row['longitud_revision']:.7f}"
        ),
        axis=1,
    )
    geographic["universo_revision"] = universe
    geographic["estado_revision_manual"] = "PENDIENTE"
    geographic["establecimiento_encontrado"] = ""
    geographic["nombre_encontrado"] = ""
    geographic["giro_observado"] = ""
    geographic["evidencia_proceso_humedo"] = ""
    geographic["evidencia_mezclilla"] = ""
    geographic["decision_manual"] = ""
    geographic["confianza_revision"] = ""
    geographic["fuentes_consultadas"] = ""
    geographic["fecha_revision"] = ""
    geographic["revisor"] = ""
    geographic["notas_revision"] = ""
    preferred = [
        "id", "clee", "nombre_de_la_unidad_economica", "raz_social",
        "codigo_de_la_clase_actividad_scian",
        "nombre_clase_actividad_scian", "categoria_clasificacion", "decision_clasificacion",
        "decision_refinada", "prioridad_refinada", "regla_refinamiento", "motivo_refinamiento",
        "per_ocu", "direccion_completa_denue", "tipo_vial", "nom_vial",
        "tipo_v_e_1", "nom_v_e_1", "tipo_v_e_2", "nom_v_e_2", "tipo_v_e_3", "nom_v_e_3",
        "numero_ext", "letra_ext", "edificio", "edificio_e", "numero_int", "letra_int",
        "tipo_asent", "nomb_asent", "cod_postal", "localidad", "municipio", "entidad",
        "ageb", "manzana", "telefono", "correoelec", "www", "tipounieco", "fecha_alta",
        "universo_revision", "latitud_revision", "longitud_revision",
        "url_busqueda_google_maps", "url_punto_google_maps", "url_openstreetmap",
        "estado_revision_manual", "establecimiento_encontrado", "nombre_encontrado",
        "giro_observado", "evidencia_proceso_humedo", "evidencia_mezclilla",
        "decision_manual", "confianza_revision", "fuentes_consultadas", "fecha_revision",
        "revisor", "notas_revision", "causa_revision_explicita",
        "decision_automatica_revision", "fuente_inclusion_automatica",
    ]
    return pd.DataFrame(geographic[[column for column in preferred if column in geographic.columns]])


def main() -> None:
    if not STAGE1_INPUT.exists() or not STAGE1_MANIFEST.exists():
        raise FileNotFoundError("Primero ejecute 02_classify_independent_universe.py")
    broad = gpd.read_file(STAGE1_INPUT, engine="pyogrio")
    policy_path = CONFIG_DIR / "politica_refinamiento.csv"
    policy = pd.read_csv(policy_path)
    OUT_DIR.joinpath("capas_qgis").mkdir(parents=True, exist_ok=True)
    OUT_DIR.joinpath("tablas").mkdir(parents=True, exist_ok=True)

    names = broad["nombre_de_la_unidad_economica"].fillna("").map(normalize_text)
    codes = broad["codigo_de_la_clase_actividad_scian"].fillna("").astype(str).str.replace(r"\.0$", "", regex=True)
    decisions = []
    for position, (_, row) in enumerate(broad.iterrows()):
        name = names.iloc[position]
        flags = {
            "is_scian_812210": codes.iloc[position] == "812210",
            "name_tintoreria": "tintoreria" in name,
            "name_industrial": any(term in name for term in ["industrial", "industria", "planta", "fabrica"]),
            "name_proceso_productivo": any(
                term in name for term in
                ["mezclilla", "denim", "stone lav", "stone wash", "lavado textil", "tenido textil", "acabado textil"]
            ),
            "original_incluir_alta": row["decision_clasificacion"] == "INCLUIR_ALTA",
        }
        selected = choose_rule(flags, policy)
        decisions.append({
            "regla_refinamiento": selected["refinement_rule"],
            "decision_refinada": selected["refined_decision"],
            "prioridad_refinada": selected["refined_priority"],
            "motivo_refinamiento": selected["reason"],
            "flags_refinamiento": "|".join(flag for flag, active in flags.items() if active),
        })
    refined = gpd.GeoDataFrame(
        pd.concat([broad.reset_index(drop=True), pd.DataFrame(decisions)], axis=1),
        geometry="geometry",
        crs=broad.crs,
    )
    included_broad = refined.loc[refined["decision_refinada"].eq("RETENER")].copy()
    review = refined.loc[refined["decision_refinada"].eq("REVISAR")].copy()
    excluded = refined.loc[refined["decision_refinada"].eq("EXCLUIR_ALCANCE")].copy()
    pending = gpd.read_file(STAGE1_REVIEW, engine="pyogrio")
    pending["_id_key"] = normalize_id(pending["id"])
    mh_ids = load_mh_ids()
    include_from_mh_mask = (
        pending["categoria_clasificacion"].eq("revisar_lavado_generico")
        & pending["_id_key"].isin(mh_ids)
    )
    pending["causa_revision_explicita"] = pending["categoria_clasificacion"].map(explicit_review_cause)
    pending["decision_automatica_revision"] = "MANTENER_EN_REVISION"
    pending["fuente_inclusion_automatica"] = ""
    pending.loc[include_from_mh_mask, "decision_automatica_revision"] = "INCLUIR_POR_MH"
    pending.loc[include_from_mh_mask, "fuente_inclusion_automatica"] = "mh_filtrado"
    accepted_mh = pending.loc[include_from_mh_mask].copy()
    accepted_mh["regla_refinamiento"] = "incluir_lavanderia_presente_mh"
    accepted_mh["decision_refinada"] = "RETENER"
    accepted_mh["prioridad_refinada"] = "media"
    accepted_mh["motivo_refinamiento"] = (
        "Lavandería genérica incluida explícitamente por pertenencia a MH; "
        "la fuente de inclusión queda trazada y no implica prioridad hídrica alta"
    )
    accepted_mh["flags_refinamiento"] = "lavanderia_generica|en_mh"
    included = gpd.GeoDataFrame(
        pd.concat([included_broad, accepted_mh], ignore_index=True),
        geometry="geometry",
        crs=refined.crs,
    )
    export(refined, "refinamiento_completo")
    export(included, "universo_productivo_refinado")
    export(review, "revision_refinamiento")
    export(excluded, "exclusiones_refinamiento")
    export(accepted_mh, "inclusiones_lavanderias_mh")
    broad_queue = manual_review_queue(refined, "universo_productivo_amplio_314")
    priority_order = {"alta": 0, "excluir": 1, "revision": 2, "media": 3}
    broad_queue["_orden"] = broad_queue.get("prioridad_refinada", "media").map(priority_order).fillna(9)
    broad_queue.sort_values(["_orden", "id"]).drop(columns="_orden").to_csv(
        OUT_DIR / "tablas" / "cola_revision_manual_314.csv", index=False, encoding="utf-8-sig"
    )
    if STAGE1_REVIEW.exists():
        review_20 = manual_review_queue(pending, "revision_fuera_universo_productivo")
        review_20.to_csv(OUT_DIR / "tablas" / "cola_revision_manual_20.csv", index=False, encoding="utf-8-sig")
        review_20.to_csv(OUT_DIR / "tablas" / "cola_reauditoria_lavanderias_20.csv", index=False, encoding="utf-8-sig")
        reaudit_xlsx = OUT_DIR / "tablas" / "cola_reauditoria_lavanderias_20.xlsx"
        try:
            review_20.to_excel(reaudit_xlsx, index=False, engine="openpyxl")
        except PermissionError:
            reaudit_xlsx = OUT_DIR / "tablas" / "cola_reauditoria_lavanderias_20_actualizada.xlsx"
            review_20.to_excel(reaudit_xlsx, index=False, engine="openpyxl")
        format_review_workbook(reaudit_xlsx)

    summary = refined.groupby(
        ["decision_refinada", "prioridad_refinada", "regla_refinamiento"], dropna=False
    ).size().reset_index(name="registros")
    summary.to_csv(OUT_DIR / "tablas" / "resumen_refinamiento.csv", index=False, encoding="utf-8-sig")
    validation = pd.DataFrame([
        {
            "validacion": "particion_conserva_universo_amplio",
            "resultado": "OK" if len(included_broad) + len(review) + len(excluded) == len(broad) else "ERROR",
        },
        {
            "validacion": "tintoreria_comercial_no_prioritaria",
            "resultado": "OK" if excluded["regla_refinamiento"].eq("excluir_tintoreria_comercial").all() else "ERROR",
        },
        {
            "validacion": "ids_unicos",
            "resultado": "OK" if not included["id"].duplicated().any() else "ERROR",
        },
        {
            "validacion": "solo_lavanderias_mh_incluidas_desde_revision",
            "resultado": "OK" if (
                accepted_mh["categoria_clasificacion"].eq("revisar_lavado_generico").all()
                and set(accepted_mh["_id_key"]).issubset(mh_ids)
            ) else "ERROR",
        },
    ])
    validation.to_csv(OUT_DIR / "tablas" / "validacion_refinamiento.csv", index=False, encoding="utf-8-sig")
    if validation["resultado"].eq("ERROR").any():
        raise RuntimeError("Falló la validación del refinamiento")
    manifest = {
        "pipeline": "refinamiento_y_priorizacion_santa_ana",
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "stage1_manifest_sha256": hashlib.sha256(STAGE1_MANIFEST.read_bytes()).hexdigest(),
        "policy_sha256": hashlib.sha256(policy_path.read_bytes()).hexdigest(),
        "input_universe": len(broad),
        "retained_from_broad": len(included_broad),
        "included_from_mh_review": len(accepted_mh),
        "retained_final_before_special_audit": len(included),
        "review": len(review),
        "excluded_scope": len(excluded),
    }
    (OUT_DIR / "manifest_refinamiento.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(summary.to_string(index=False))


if __name__ == "__main__":
    main()
