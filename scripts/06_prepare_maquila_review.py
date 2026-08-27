from __future__ import annotations

from pathlib import Path
from urllib.parse import quote_plus

import geopandas as gpd
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from common import PROJECT_ROOT, normalize_text


INPUT_PATH = (
    PROJECT_ROOT / "outputs" / "auditoria_6_especiales" / "capas_qgis"
    / "universo_productivo_auditado.gpkg"
)
OUT_DIR = PROJECT_ROOT / "outputs" / "revision_maquila"
MH_PATH = PROJECT_ROOT / "data" / "processed" / "capa_maryhelen" / "capa MH.shp"


def maps_search(row: pd.Series) -> str:
    parts = [str(row.get("nombre_de_la_unidad_economica", ""))]
    for column in ["nom_vial", "nomb_asent", "cod_postal"]:
        value = str(row.get(column, "") or "").strip()
        if value and value.lower() != "nan":
            parts.append(value)
    parts.extend(["Santa Ana Xalmimilulco", "Puebla", "México"])
    return "https://www.google.com/maps/search/?api=1&query=" + quote_plus(" ".join(parts))


def clean_value(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def full_address(row: pd.Series) -> str:
    street = " ".join(
        value for value in [
            clean_value(row.get("tipo_vial")),
            clean_value(row.get("nom_vial")),
            clean_value(row.get("numero_ext")),
            clean_value(row.get("letra_ext")),
        ] if value
    )
    settlement = " ".join(
        value for value in [
            clean_value(row.get("tipo_asent")),
            clean_value(row.get("nomb_asent")),
        ] if value
    )
    return ", ".join(
        value for value in [
            street,
            settlement,
            clean_value(row.get("cod_postal")),
            clean_value(row.get("localidad")),
            clean_value(row.get("municipio")),
            clean_value(row.get("entidad")),
        ] if value
    )


def format_review_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    link_columns = {
        "url_busqueda_google_maps": "Abrir búsqueda en Google Maps",
        "url_punto_google_maps": "Abrir punto en Google Maps",
        "url_openstreetmap": "Abrir punto en OpenStreetMap",
    }
    for column_name, label in link_columns.items():
        column = headers[column_name]
        for row_number in range(2, sheet.max_row + 1):
            cell = sheet.cell(row_number, column)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.value = label
                cell.font = Font(color="0563C1", underline="single")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def main() -> None:
    if not INPUT_PATH.exists():
        raise FileNotFoundError("Primero ejecute 05_apply_special_audit.py")
    universe = gpd.read_file(INPUT_PATH, engine="pyogrio").to_crs("EPSG:4326")
    maquila = universe.loc[universe["categoria_clasificacion"].eq("maquila_textil")].copy()
    mh = gpd.read_file(MH_PATH, engine="pyogrio")
    mh_ids = set(
        mh["ID DENUE"].fillna("").astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
    ) - {"", "nan", "None"}
    maquila["en_mh"] = maquila["id"].astype(str).str.replace(r"\.0$", "", regex=True).isin(mh_ids)
    maquila["antecedente_mh"] = "NO_APARECE_EN_MH"
    maquila.loc[maquila["en_mh"], "antecedente_mh"] = "VISTO_BUENO_PREVIO_MH"
    maquila["nota_antecedente_mh"] = ""
    maquila.loc[maquila["en_mh"], "nota_antecedente_mh"] = (
        "MH incluyó previamente este ID DENUE; usar como antecedente favorable, "
        "sin asumir que confirma proceso húmedo actual"
    )
    maquila["latitud_revision"] = maquila.geometry.y
    maquila["longitud_revision"] = maquila.geometry.x
    maquila["direccion_completa"] = maquila.apply(full_address, axis=1)
    maquila["url_busqueda_google_maps"] = maquila.apply(maps_search, axis=1)
    maquila["url_punto_google_maps"] = maquila.apply(
        lambda row: "https://www.google.com/maps/search/?api=1&query="
        + quote_plus(f"{row['latitud_revision']},{row['longitud_revision']}"), axis=1
    )
    maquila["url_openstreetmap"] = maquila.apply(
        lambda row: (
            f"https://www.openstreetmap.org/?mlat={row['latitud_revision']:.7f}"
            f"&mlon={row['longitud_revision']:.7f}#map=19/"
            f"{row['latitud_revision']:.7f}/{row['longitud_revision']:.7f}"
        ), axis=1
    )
    names = maquila["nombre_de_la_unidad_economica"].fillna("").map(normalize_text)
    wet_pattern = (
        r"lavado|lavander|deslav|deslave|stone wash|stone lav|acid wash|tenid|tintura|"
        r"blanque|mercer|enzim|sosa caustica|colorante|bano quimico|quimic|"
        r"suaviz|pigment|decolor|neutraliz|centrifug|enjuag|oxidante"
    )
    denim_pattern = r"mezclilla|mesclilla|denim|jeans?"
    dry_pattern = r"\bcorte\b|costur|confeccion|deshebr|bordad|\bover\b|\brecta\b|hojale|ensambl"
    indirect_dry_pattern = r"planch"
    ambiguous_finish_pattern = r"acabad|terminado"

    wet_mask = names.str.contains(wet_pattern, regex=True)
    denim_mask = names.str.contains(denim_pattern, regex=True)
    dry_mask = names.str.contains(dry_pattern, regex=True)
    indirect_dry_mask = names.str.contains(indirect_dry_pattern, regex=True)
    ambiguous_finish_mask = names.str.contains(ambiguous_finish_pattern, regex=True)

    maquila["clasificacion_proceso_hidrico"] = "MAQUILA_SIN_PROCESO_DECLARADO"
    maquila["decision_hidrica_automatica"] = "REVISAR"
    maquila["prioridad_hidrica_maquila"] = "revision"
    maquila["causa_decision_hidrica"] = (
        "DENUE confirma maquila/confección SCIAN 315, pero el nombre no declara el proceso; "
        "se requiere auditoría para descartar operaciones húmedas"
    )

    # La precedencia metodológica es húmedo > mezclilla > seco > acabado ambiguo > genérico.
    maquila.loc[ambiguous_finish_mask, "clasificacion_proceso_hidrico"] = "ACABADO_AMBIGUO"
    maquila.loc[ambiguous_finish_mask, "causa_decision_hidrica"] = (
        "El nombre declara acabado o terminado sin especificar si utiliza agua o químicos"
    )
    maquila.loc[dry_mask | indirect_dry_mask, "clasificacion_proceso_hidrico"] = "SECO_EXPLICITO"
    maquila.loc[dry_mask | indirect_dry_mask, "decision_hidrica_automatica"] = "EXCLUIR_SECO"
    maquila.loc[dry_mask | indirect_dry_mask, "prioridad_hidrica_maquila"] = "fuera_hidrico"
    maquila.loc[dry_mask, "causa_decision_hidrica"] = (
        "El nombre declara corte, costura, confección, deshebrado, bordado o ensamble; "
        "proceso seco fuera del universo hídrico"
    )
    maquila.loc[indirect_dry_mask, "causa_decision_hidrica"] = (
        "El nombre declara planchado; puede usar vapor indirectamente, pero no constituye "
        "un proceso húmedo textil prioritario"
    )
    maquila.loc[denim_mask, "clasificacion_proceso_hidrico"] = "MEZCLILLA_DENIM_EXPLICITA"
    maquila.loc[denim_mask, "decision_hidrica_automatica"] = "INCLUIR_POR_MEZCLILLA"
    maquila.loc[denim_mask, "prioridad_hidrica_maquila"] = "inclusion_precautoria"
    maquila.loc[denim_mask, "causa_decision_hidrica"] = (
        "El nombre DENUE menciona explícitamente mezclilla, mesclilla, denim o jeans; "
        "se incluye por política precautoria, aunque el nombre no pruebe lavado"
    )
    maquila.loc[wet_mask, "clasificacion_proceso_hidrico"] = "HUMEDO_EXPLICITO"
    maquila.loc[wet_mask, "decision_hidrica_automatica"] = "INCLUIR_PROCESO_HUMEDO"
    maquila.loc[wet_mask, "prioridad_hidrica_maquila"] = "alta"
    maquila.loc[wet_mask, "causa_decision_hidrica"] = (
        "El nombre DENUE declara una señal fuerte de lavado, teñido, blanqueo, "
        "mercerizado o acabado químico/húmedo"
    )

    maquila["grupo_primera_revision"] = "maquila_generica_scian_315"
    maquila.loc[names.str.contains("pantal", regex=False), "grupo_primera_revision"] = "pantalon_sin_mezclilla_explicita"
    maquila.loc[names.str.contains(r"falda|camisa|uniforme|ropa|vestir", regex=True), "grupo_primera_revision"] = "otra_prenda_explicita"
    maquila.loc[ambiguous_finish_mask, "grupo_primera_revision"] = "acabado_ambiguo"
    maquila.loc[dry_mask | indirect_dry_mask, "grupo_primera_revision"] = "proceso_seco_explicito"
    maquila.loc[denim_mask, "grupo_primera_revision"] = "mezclilla_explicita"
    maquila.loc[wet_mask, "grupo_primera_revision"] = "proceso_humedo_explicito"

    projected = maquila.to_crs("EPSG:6372")
    geometry_key = projected.geometry.to_wkb(hex=True)
    possible_duplicate = geometry_key.duplicated(keep=False)
    maquila["posible_duplicado_geometria"] = possible_duplicate.to_numpy()
    maquila["ids_misma_geometria"] = ""
    for _, group in maquila.loc[possible_duplicate].groupby(geometry_key[possible_duplicate].to_numpy()):
        ids = "|".join(group["id"].astype(str))
        maquila.loc[group.index, "ids_misma_geometria"] = ids

    maquila["dictamen_documental_inicial"] = "GIRO_TEXTIL_RESPALDADO_POR_SCIAN_315"
    maquila["error_sectorial_evidente"] = False
    maquila["causa_dictamen_inicial"] = (
        "Código SCIAN manufacturero 315 y descripción de confección de prendas; "
        "confirma pertenencia al universo productivo, no existencia actual ni proceso húmedo"
    )
    priority = {
        "proceso_humedo_explicito": 0,
        "mezclilla_explicita": 1,
        "acabado_ambiguo": 2,
        "pantalon_sin_mezclilla_explicita": 3,
        "otra_prenda_explicita": 3,
        "maquila_generica_scian_315": 4,
        "proceso_seco_explicito": 9,
    }
    maquila["orden_revision"] = maquila["grupo_primera_revision"].map(priority).fillna(9).astype(int)
    maquila.loc[possible_duplicate, "orden_revision"] = 0
    maquila["estado_revision_manual"] = "PENDIENTE"
    maquila.loc[
        maquila["decision_hidrica_automatica"].ne("REVISAR"),
        "estado_revision_manual",
    ] = "RESUELTO_AUTOMATICAMENTE"
    maquila["establecimiento_encontrado"] = ""
    maquila["objeto_maquila_observado"] = ""
    maquila["evidencia_textil"] = ""
    maquila["evidencia_mezclilla"] = ""
    maquila["evidencia_proceso_humedo"] = ""
    maquila["decision_manual"] = ""
    maquila["confianza_revision"] = ""
    maquila["fuentes_consultadas"] = ""
    maquila["fecha_revision"] = ""
    maquila["revisor"] = ""
    maquila["notas_revision"] = ""
    columns = [
        "id", "clee", "nombre_de_la_unidad_economica", "raz_social",
        "codigo_de_la_clase_actividad_scian",
        "nombre_clase_actividad_scian", "categoria_clasificacion", "regla_decisiva",
        "terminos_detectados", "per_ocu", "direccion_completa",
        "tipo_vial", "nom_vial", "numero_ext", "letra_ext", "numero_int", "letra_int",
        "tipo_asent", "nomb_asent", "cod_postal", "localidad", "municipio", "entidad",
        "telefono", "correoelec", "www", "latitud_revision", "longitud_revision",
        "url_busqueda_google_maps", "url_punto_google_maps", "url_openstreetmap",
        "clasificacion_proceso_hidrico", "decision_hidrica_automatica",
        "prioridad_hidrica_maquila", "causa_decision_hidrica",
        "en_mh", "antecedente_mh", "nota_antecedente_mh",
        "grupo_primera_revision", "orden_revision", "dictamen_documental_inicial",
        "error_sectorial_evidente", "causa_dictamen_inicial",
        "posible_duplicado_geometria", "ids_misma_geometria",
        "estado_revision_manual", "establecimiento_encontrado", "objeto_maquila_observado",
        "evidencia_textil", "evidencia_mezclilla", "evidencia_proceso_humedo",
        "decision_manual", "confianza_revision", "fuentes_consultadas", "fecha_revision",
        "revisor", "notas_revision",
    ]
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    queue = maquila[columns].sort_values(["orden_revision", "nombre_de_la_unidad_economica", "id"])
    xlsx_path = OUT_DIR / "cola_revision_maquila.xlsx"
    try:
        queue.to_excel(xlsx_path, index=False, engine="openpyxl")
    except PermissionError:
        xlsx_path = OUT_DIR / "cola_revision_maquila_actualizada.xlsx"
        queue.to_excel(xlsx_path, index=False, engine="openpyxl")
        print(
            "Aviso: cola_revision_maquila.xlsx está abierto; "
            "se escribió cola_revision_maquila_actualizada.xlsx"
        )
    format_review_workbook(xlsx_path)
    csv_path = OUT_DIR / "cola_revision_maquila.csv"
    try:
        queue.to_csv(csv_path, index=False, encoding="utf-8-sig")
    except PermissionError:
        csv_path = OUT_DIR / "cola_revision_maquila_actualizada.csv"
        queue.to_csv(csv_path, index=False, encoding="utf-8-sig")
        print(
            "Aviso: cola_revision_maquila.csv está abierto; "
            "se escribió cola_revision_maquila_actualizada.csv"
        )
    summary = (
        maquila.groupby(
            ["clasificacion_proceso_hidrico", "decision_hidrica_automatica"], dropna=False
        )
        .size().reset_index(name="registros").sort_values("registros", ascending=False)
    )
    summary.to_csv(OUT_DIR / "resumen_primera_revision_maquila.csv", index=False, encoding="utf-8-sig")
    queue.loc[queue["posible_duplicado_geometria"]].to_csv(
        OUT_DIR / "posibles_duplicados_maquila.csv", index=False, encoding="utf-8-sig"
    )
    queue.loc[queue["decision_hidrica_automatica"].str.startswith("INCLUIR")].to_csv(
        OUT_DIR / "maquilas_inclusion_automatica_hidrica.csv", index=False, encoding="utf-8-sig"
    )
    queue.loc[queue["decision_hidrica_automatica"].eq("EXCLUIR_SECO")].to_csv(
        OUT_DIR / "maquilas_excluidas_proceso_seco.csv", index=False, encoding="utf-8-sig"
    )
    pending_review = queue.loc[queue["decision_hidrica_automatica"].eq("REVISAR")]
    pending_xlsx = OUT_DIR / "cola_revision_maquila_pendiente.xlsx"
    try:
        pending_review.to_excel(pending_xlsx, index=False, engine="openpyxl")
    except PermissionError:
        pending_xlsx = OUT_DIR / "cola_revision_maquila_pendiente_actualizada.xlsx"
        pending_review.to_excel(pending_xlsx, index=False, engine="openpyxl")
    format_review_workbook(pending_xlsx)
    try:
        pending_review.to_csv(
            OUT_DIR / "cola_revision_maquila_pendiente.csv", index=False, encoding="utf-8-sig"
        )
    except PermissionError:
        pending_review.to_csv(
            OUT_DIR / "cola_revision_maquila_pendiente_actualizada.csv",
            index=False, encoding="utf-8-sig",
        )
    print(f"Cola de maquila preparada: {len(maquila)} registros")


if __name__ == "__main__":
    main()
